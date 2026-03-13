"""
Load care intent rules from JSON, Excel (.xlsx), or CSV.
Edit rules in Word/Excel/Sheets on your machine; code loads the file and uses it for matching + training.
"""
import csv
import json
import logging
from pathlib import Path
from typing import Any, Dict, List, Tuple

logger = logging.getLogger(__name__)


def _split_cell(val: Any) -> List[str]:
    if val is None or (isinstance(val, float) and (val != val or val == 0)):
        return []
    s = str(val).strip()
    if not s:
        return []
    # Allow commas, semicolons, or pipes as separators
    for sep in (";", "|"):
        s = s.replace(sep, ",")
    return [p.strip() for p in s.split(",") if p.strip()]


def _bool_cell(val: Any) -> bool:
    if val is None:
        return False
    if isinstance(val, bool):
        return val
    s = str(val).strip().lower()
    return s in ("1", "true", "yes", "y")


def _num_cell(val: Any) -> int | None:
    if val is None or val == "":
        return None
    try:
        return int(float(val))
    except (TypeError, ValueError):
        return None


def load_from_json(path: Path) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    intents = list(data.get("intents") or [])
    out_of_scope = data.get("out_of_scope") or {}
    return intents, out_of_scope


def load_from_xlsx(path: Path) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    try:
        import openpyxl
    except ImportError:
        logger.warning("care_rules: openpyxl not installed — pip install openpyxl in the venv used to run the app")
        return [], {}

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    intents: List[Dict[str, Any]] = []
    out_of_scope: Dict[str, Any] = {}

    if "out_of_scope" in wb.sheetnames:
        ws = wb["out_of_scope"]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) >= 2:
            headers = [str(c or "").strip().lower().replace(" ", "_") for c in rows[0]]
            for row in rows[1:]:
                if not any(row):
                    continue
                d = dict(zip(headers, (c for c in row)))
                if d.get("response"):
                    out_of_scope["response"] = str(d["response"]).strip()
                if d.get("suggested_actions"):
                    out_of_scope["suggested_actions"] = _split_cell(d["suggested_actions"])
                break

    sheet_names = wb.sheetnames
    if "intents" in sheet_names:
        ws = wb["intents"]
    elif "Training Data" in sheet_names:
        ws = wb["Training Data"]
    else:
        ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        logger.debug("care_rules: %s sheet has < 2 rows", path.name)
        wb.close()
        return intents, out_of_scope

    headers = [str(c or "").strip().lower().replace(" ", "_") for c in rows[0]]

    # Training-data style: Intent + phrase column (one row per phrase)
    lower_headers = [h.lower() for h in headers]
    has_intent = "intent" in lower_headers or "id" in lower_headers
    idx_phrase_candidate = next(
        (i for i, h in enumerate(lower_headers) if h.startswith("user_input")),
        None,
    )
    if idx_phrase_candidate is None:
        idx_phrase_candidate = next(
            (i for i, h in enumerate(lower_headers) if "phrase" in h or "user_input" in h),
            None,
        )
    if has_intent and idx_phrase_candidate is not None:
        idx_intent = lower_headers.index("intent") if "intent" in lower_headers else lower_headers.index("id")
        idx_response = next((i for i, h in enumerate(lower_headers) if "expected_response" in h or h == "response"), None)
        idx_short = next((i for i, h in enumerate(lower_headers) if "short_response" in h), None)
        idx_suggested = next((i for i, h in enumerate(lower_headers) if "suggested_actions" in h), None)
        idx_phrase = idx_phrase_candidate

        by_intent: Dict[str, Dict[str, Any]] = {}
        for row in rows[1:]:
            if not any(row):
                continue
            intent_raw = row[idx_intent] if idx_intent is not None and idx_intent < len(row) else None
            phrase_raw = row[idx_phrase] if idx_phrase is not None and idx_phrase < len(row) else None
            if not intent_raw or not phrase_raw:
                continue
            intent_id = str(intent_raw).strip()
            phrase = str(phrase_raw).strip()
            if not intent_id or not phrase:
                continue
            r = by_intent.setdefault(intent_id, {"id": intent_id, "phrases": []})
            r["phrases"].append(phrase)
            if idx_response is not None and not r.get("response") and idx_response < len(row) and row[idx_response]:
                r["response"] = str(row[idx_response]).strip()
            if idx_short is not None and not r.get("short_response") and idx_short < len(row) and row[idx_short]:
                r["short_response"] = str(row[idx_short]).strip()
            if idx_suggested is not None and not r.get("suggested_actions") and idx_suggested < len(row) and row[idx_suggested]:
                r["suggested_actions"] = _split_cell(row[idx_suggested])
        intents = list(by_intent.values())
    else:
        # No standard headers: try first row as data (col0=intent, col1=phrase, col2=response)
        first_cell = str((rows[0][0]) if rows[0] else "").strip()
        if len(rows[0]) >= 2 and first_cell and ("_" in first_cell or first_cell.isupper() or len(first_cell) <= 30):
            by_intent = {}
            for row in rows:
                if not row or len(row) < 2:
                    continue
                intent_id = str(row[0] or "").strip()
                phrase = str(row[1] or "").strip()
                if not intent_id or not phrase:
                    continue
                r = by_intent.setdefault(intent_id, {"id": intent_id, "phrases": []})
                r["phrases"].append(phrase)
                if len(row) >= 3 and row[2]:
                    r.setdefault("response", str(row[2]).strip())
            intents = list(by_intent.values())
        else:
            for row in rows[1:]:
                if not any(row):
                    continue
                d = dict(zip(headers, row))
                intent_id = (d.get("intent_id") or d.get("id") or "").strip()
                if not intent_id or intent_id.upper() == "OUT_OF_SCOPE":
                    if d.get("response"):
                        out_of_scope["response"] = str(d.get("response") or "").strip()
                    if d.get("suggested_actions"):
                        out_of_scope["suggested_actions"] = _split_cell(d["suggested_actions"])
                    continue

                r: Dict[str, Any] = {"id": intent_id}
                if d.get("phrases"):
                    r["phrases"] = _split_cell(d["phrases"])
                if d.get("response"):
                    r["response"] = str(d["response"]).strip()
                if d.get("short_response"):
                    r["short_response"] = str(d["short_response"]).strip()
                if d.get("short_phrases"):
                    r["short_phrases"] = _split_cell(d["short_phrases"])
                if d.get("suggested_actions"):
                    r["suggested_actions"] = _split_cell(d["suggested_actions"])
                if d.get("response_no_contact"):
                    r["response_no_contact"] = str(d["response_no_contact"]).strip()
                if d.get("response_with_contact"):
                    r["response_with_contact"] = str(d["response_with_contact"]).strip()
                r["prepend_tx_context"] = _bool_cell(d.get("prepend_tx_context"))
                r["append_kb"] = _bool_cell(d.get("append_kb"))
                nm = _num_cell(d.get("max_message_len"))
                if nm is not None:
                    r["max_message_len"] = nm
                intents.append(r)
    wb.close()
    return intents, out_of_scope


def load_from_csv(path: Path) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    intents: List[Dict[str, Any]] = []
    out_of_scope: Dict[str, Any] = {}
    with open(path, encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            intent_id = (row.get("intent_id") or row.get("id") or "").strip()
            if not intent_id or intent_id.upper() == "OUT_OF_SCOPE":
                if row.get("response"):
                    out_of_scope["response"] = row["response"].strip()
                if row.get("suggested_actions"):
                    out_of_scope["suggested_actions"] = _split_cell(row["suggested_actions"])
                continue
            r: Dict[str, Any] = {"id": intent_id}
            if row.get("phrases"):
                r["phrases"] = _split_cell(row["phrases"])
            if row.get("response"):
                r["response"] = row["response"].strip()
            if row.get("short_response"):
                r["short_response"] = row["short_response"].strip()
            if row.get("short_phrases"):
                r["short_phrases"] = _split_cell(row["short_phrases"])
            if row.get("suggested_actions"):
                r["suggested_actions"] = _split_cell(row["suggested_actions"])
            if row.get("response_no_contact"):
                r["response_no_contact"] = row["response_no_contact"].strip()
            if row.get("response_with_contact"):
                r["response_with_contact"] = row["response_with_contact"].strip()
            r["prepend_tx_context"] = _bool_cell(row.get("prepend_tx_context"))
            r["append_kb"] = _bool_cell(row.get("append_kb"))
            nm = _num_cell(row.get("max_message_len"))
            if nm is not None:
                r["max_message_len"] = nm
            intents.append(r)
    return intents, out_of_scope


def load_from_docx(path: Path) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """Read first table = intents (header row), second table = out_of_scope (optional)."""
    try:
        from docx import Document
    except ImportError:
        return [], {}

    doc = Document(path)
    intents = []
    out_of_scope = {}
    tables = doc.tables
    if not tables:
        return [], {}

    headers = [str(c.text or "").strip().lower().replace(" ", "_") for c in tables[0].rows[0].cells]
    for row in tables[0].rows[1:]:
        cells = [c.text for c in row.cells]
        d = dict(zip(headers, (cells[i] if i < len(cells) else None for i in range(len(headers)))))
        intent_id = (d.get("intent_id") or d.get("id") or "").strip()
        if not intent_id or intent_id.upper() == "OUT_OF_SCOPE":
            if d.get("response"):
                out_of_scope["response"] = str(d.get("response") or "").strip()
            if d.get("suggested_actions"):
                out_of_scope["suggested_actions"] = _split_cell(d["suggested_actions"])
            continue
        r: Dict[str, Any] = {"id": intent_id}
        if d.get("phrases"):
            r["phrases"] = _split_cell(d["phrases"])
        if d.get("response"):
            r["response"] = str(d["response"]).strip()
        if d.get("short_response"):
            r["short_response"] = str(d["short_response"]).strip()
        if d.get("short_phrases"):
            r["short_phrases"] = _split_cell(d["short_phrases"])
        if d.get("suggested_actions"):
            r["suggested_actions"] = _split_cell(d["suggested_actions"])
        if d.get("response_no_contact"):
            r["response_no_contact"] = str(d["response_no_contact"]).strip()
        if d.get("response_with_contact"):
            r["response_with_contact"] = str(d["response_with_contact"]).strip()
        r["prepend_tx_context"] = _bool_cell(d.get("prepend_tx_context"))
        r["append_kb"] = _bool_cell(d.get("append_kb"))
        nm = _num_cell(d.get("max_message_len"))
        if nm is not None:
            r["max_message_len"] = nm
        intents.append(r)

    if len(tables) >= 2:
        h2 = [str(c.text or "").strip().lower().replace(" ", "_") for c in tables[1].rows[0].cells]
        for row in tables[1].rows[1:]:
            cells = [c.text for c in row.cells]
            d = dict(zip(h2, (cells[i] if i < len(cells) else None for i in range(len(h2)))))
            if d.get("response"):
                out_of_scope["response"] = str(d["response"]).strip()
            if d.get("suggested_actions"):
                out_of_scope["suggested_actions"] = _split_cell(d["suggested_actions"])
            break
    return intents, out_of_scope


def load_intent_rules(data_dir: Path | None = None) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """Load rules from disk.

    Priority:
    1) Explicit bank_chatbot Excel files if present (user-authored training sheet)
    2) care_intent_rules.json
    3) care_intent_rules.xlsx
    4) care_intent_rules.docx
    5) care_intent_rules.csv
    """
    if data_dir is None:
        data_dir = Path(__file__).resolve().parent / "data"

    # 1) Prefer explicit bank_chatbot Excel files if the user has dropped one in ml/data,
    #    but only if they actually contain intents; otherwise continue to JSON.
    for fname in ("bank_chatbot_v2.xlsx", "bank_chatbbot.xlsx", "bank_chatbot_training_data.xlsx"):
        alt = data_dir / fname
        if alt.exists():
            intents, out_of_scope = load_from_xlsx(alt)
            if intents:
                logger.info("care_rules: loaded %d intents from %s", len(intents), fname)
                print(f"[CARE_RULES] loaded {len(intents)} intents from {fname}")
                return intents, out_of_scope
            logger.warning("care_rules: %s returned 0 intents (check sheet 'intents', headers Intent / User Input (Phrase), or pip install openpyxl in app venv)", fname)
            print(f"[CARE_RULES] {fname} returned 0 intents — check sheet 'intents', headers Intent, User Input (Phrase); or run: pip install openpyxl")

    # 2) Fallback to the standard care_intent_rules.* files
    for ext, loader in [
        (".json", load_from_json),
        (".xlsx", load_from_xlsx),
        (".docx", load_from_docx),
        (".csv", load_from_csv),
    ]:
        path = data_dir / f"care_intent_rules{ext}"
        if path.exists():
            intents, out_of_scope = loader(path)
            logger.info("care_rules: loaded %d intents from care_intent_rules%s", len(intents), ext)
            print(f"[CARE_RULES] loaded {len(intents)} intents from care_intent_rules{ext}")
            return intents, out_of_scope
    return [], {}
